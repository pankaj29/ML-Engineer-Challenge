"""Render the deliverables checklist to an Excel workbook.

Usage::

    python scripts/generate_checklist.py

Reads :mod:`scripts.checklist_data` and writes
``DELIVERABLES_CHECKLIST.xlsx`` in the repository root. Re-run it any time a
status changes; the file is rewritten from scratch, so the spreadsheet and the
source list can never disagree.

The workbook has two sheets:
    Checklist - every requirement from README.md, with a tick box and status.
    Summary   - completion percentages per part, so progress is visible at a glance.
"""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from scripts.checklist_data import Item, apply_updates

ITEMS = apply_updates()

REPO_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = REPO_ROOT / "DELIVERABLES_CHECKLIST.xlsx"

# --- Visual style ----------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
PART_FILL = PatternFill("solid", fgColor="D9E2F3")
PART_FONT = Font(bold=True, size=11, color="1F3864")

STATUS_STYLE: dict[str, tuple[str, str]] = {
    # status -> (background colour, font colour)
    "DONE": ("C6EFCE", "006100"),
    "IN PROGRESS": ("FFEB9C", "9C5700"),
    "TODO": ("F2F2F2", "595959"),
    "BLOCKED": ("FFC7CE", "9C0006"),
}
TICK = {"DONE": "✔", "IN PROGRESS": "◐", "TODO": "☐", "BLOCKED": "✗"}

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_header(ws: Worksheet, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 28


def _build_checklist_sheet(ws: Worksheet, items: list[Item]) -> None:
    """Write the main checklist, grouped by part with a banner per group."""
    _style_header(
        ws,
        [
            "#",
            "Done",
            "Part",
            "Requirement (from README.md)",
            "Status",
            "Evidence",
            "Notes / Assumptions",
        ],
    )

    row = 2
    number = 0
    last_part_group = None

    for item in items:
        part_group = item.part.split(":")[0].strip()
        if part_group != last_part_group:
            # Group banner row, merged across the full width.
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
            banner = ws.cell(row=row, column=1, value=part_group.upper())
            banner.fill = PART_FILL
            banner.font = PART_FONT
            banner.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws.row_dimensions[row].height = 20
            row += 1
            last_part_group = part_group

        number += 1
        bg, fg = STATUS_STYLE.get(item.status, STATUS_STYLE["TODO"])
        values = [
            number,
            TICK.get(item.status, "☐"),
            item.part,
            item.requirement,
            item.status,
            item.evidence,
            item.notes,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=col in (4, 6, 7),
                horizontal="center" if col in (1, 2, 5) else "left",
            )
            if col in (2, 5):
                cell.fill = PatternFill("solid", fgColor=bg)
                cell.font = Font(color=fg, bold=True, size=12 if col == 2 else 10)
        row += 1

    # Column widths tuned so the sheet is readable without manual resizing.
    for col, width in zip("ABCDEFG", [5, 7, 20, 58, 14, 34, 52], strict=False):
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{row - 1}"


def _build_summary_sheet(ws: Worksheet, items: list[Item]) -> None:
    """Write per-part completion statistics."""
    _style_header(ws, ["Part", "Total", "Done", "In Progress", "To Do", "Blocked", "% Complete"])

    groups: dict[str, list[Item]] = {}
    for item in items:
        groups.setdefault(item.part.split(":")[0].strip(), []).append(item)

    row = 2
    for group, group_items in groups.items():
        total = len(group_items)
        done = sum(1 for i in group_items if i.status == "DONE")
        prog = sum(1 for i in group_items if i.status == "IN PROGRESS")
        todo = sum(1 for i in group_items if i.status == "TODO")
        blocked = sum(1 for i in group_items if i.status == "BLOCKED")
        pct = done / total if total else 0.0

        for col, value in enumerate([group, total, done, prog, todo, blocked, pct], start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="left" if col == 1 else "center")
            if col == 7:
                cell.number_format = "0%"
                cell.font = Font(bold=True)
        row += 1

    # Grand total row.
    total = len(items)
    done = sum(1 for i in items if i.status == "DONE")
    prog = sum(1 for i in items if i.status == "IN PROGRESS")
    todo = sum(1 for i in items if i.status == "TODO")
    blocked = sum(1 for i in items if i.status == "BLOCKED")
    for col, value in enumerate(
        ["TOTAL", total, done, prog, todo, blocked, done / total if total else 0.0], start=1
    ):
        cell = ws.cell(row=row, column=col, value=value)
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="left" if col == 1 else "center")
        if col == 7:
            cell.number_format = "0%"

    row += 2
    ws.cell(row=row, column=1, value="Generated").font = Font(bold=True)
    # Local time on purpose: this is a human-facing "generated at" stamp
    # in a spreadsheet, not a machine-comparable timestamp.
    ws.cell(
        row=row,
        column=2,
        value=datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %Z"),
    )
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    row += 1
    ws.cell(row=row, column=1, value="Source").font = Font(bold=True)
    ws.cell(
        row=row,
        column=2,
        value="scripts/checklist_data.py (regenerate with scripts/generate_checklist.py)",
    )
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)

    for col, width in zip("ABCDEFG", [26, 10, 10, 14, 10, 10, 14], strict=False):
        ws.column_dimensions[col].width = width


def generate(output: Path = OUTPUT_PATH) -> Path:
    """Build the workbook and write it to ``output``."""
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    _build_summary_sheet(summary, ITEMS)

    checklist = wb.create_sheet("Checklist")
    _build_checklist_sheet(checklist, ITEMS)

    wb.save(output)
    return output


def main() -> int:
    path = generate()
    total = len(ITEMS)
    done = sum(1 for i in ITEMS if i.status == "DONE")
    prog = sum(1 for i in ITEMS if i.status == "IN PROGRESS")
    blocked = sum(1 for i in ITEMS if i.status == "BLOCKED")
    print(f"Wrote {path}")
    print(
        f"  {done}/{total} done ({done / total:.0%}), "
        f"{prog} in progress, {blocked} blocked, {total - done - prog - blocked} to do"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
